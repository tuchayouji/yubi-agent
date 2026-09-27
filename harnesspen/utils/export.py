"""文件导出模块：支持 PDF, DOCX, MD 格式。正确处理中文编码。"""

import os
import re
from io import BytesIO

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm, mm
from reportlab.lib.enums import TA_LEFT, TA_CENTER
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Preformatted,
    ListFlowable, ListItem, PageBreak
)
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont


# ─── 字体注册 ───

_FONTS_REGISTERED = False
_CN_FONT = "MicrosoftYaHei"
_CN_FONT_BOLD = "MicrosoftYaHei-Bold"
_CN_MONO = "Courier"


def _register_fonts():
    """注册中文字体到 reportlab。"""
    global _FONTS_REGISTERED, _CN_FONT_BOLD
    if _FONTS_REGISTERED:
        return

    fonts_dir = r"C:\Windows\Fonts"

    # 尝试多种中文字体
    candidates = [
        ("msyh.ttc", _CN_FONT, 0),
        ("msyhbd.ttc", _CN_FONT_BOLD, 0),
        ("simhei.ttf", _CN_FONT, None),
        ("simfang.ttf", _CN_FONT, None),
        ("Deng.ttf", _CN_FONT, None),
    ]

    regular_registered = False
    bold_registered = False

    for filename, font_name, subfont_index in candidates:
        path = os.path.join(fonts_dir, filename)
        if not os.path.exists(path):
            continue
        try:
            if font_name == _CN_FONT and not regular_registered:
                if subfont_index is not None:
                    pdfmetrics.registerFont(TTFont(font_name, path, subfontIndex=subfont_index))
                else:
                    pdfmetrics.registerFont(TTFont(font_name, path))
                regular_registered = True
            elif font_name == _CN_FONT_BOLD and not bold_registered:
                if subfont_index is not None:
                    pdfmetrics.registerFont(TTFont(font_name, path, subfontIndex=subfont_index))
                else:
                    pdfmetrics.registerFont(TTFont(font_name, path))
                bold_registered = True
        except Exception:
            continue

    # 如果没有找到粗体，用普通体代替
    if not bold_registered and regular_registered:
        _CN_FONT_BOLD = _CN_FONT

    _FONTS_REGISTERED = True


# ─── 公共 ───

def _get_filename(article_id: str, fmt: str) -> str:
    """生成下载文件名。"""
    return f"{article_id}.{fmt}"


# ─── MD ───

def export_md(content: str) -> bytes:
    """直接返回 UTF-8 编码的 Markdown 内容。"""
    return content.encode("utf-8")


# ─── DOCX ───

def _set_chinese_font(run, font_name="微软雅黑"):
    """为 run 设置中文字体（包含 East Asian 属性）。"""
    run.font.name = font_name
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = r.makeelement(qn('w:rFonts'), {})
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:eastAsia'), font_name)


def _set_heading_chinese(paragraph, font_name="微软雅黑"):
    """为标题段落设置中文字体。"""
    for run in paragraph.runs:
        _set_chinese_font(run, font_name)


def export_docx(content: str) -> bytes:
    """将 Markdown 转换为 .docx 文件，正确处理中文。"""
    doc = Document()

    # 设置默认字体
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(11)
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_after = Pt(6)
    # 设置 East Asian 字体
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = style.element.makeelement(qn('w:rFonts'), {})
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:eastAsia'), '微软雅黑')

    # 设置标题样式
    for i in range(1, 4):
        hs = doc.styles[f"Heading {i}"]
        hs.font.name = "微软雅黑"
        hs.font.color.rgb = RGBColor(0x31, 0x2E, 0x81)
        rPr = hs.element.get_or_add_rPr()
        rFonts = rPr.find(qn('w:rFonts'))
        if rFonts is None:
            rFonts = hs.element.makeelement(qn('w:rFonts'), {})
            rPr.insert(0, rFonts)
        rFonts.set(qn('w:eastAsia'), '微软雅黑')

    lines = content.split("\n")
    in_code_block = False
    code_buffer = []

    for line in lines:
        # 代码块开关
        if line.startswith("```"):
            if in_code_block:
                code_text = "\n".join(code_buffer)
                p = doc.add_paragraph()
                p.style = doc.styles["Normal"]
                p.paragraph_format.left_indent = Inches(0.3)
                run = p.add_run(code_text)
                run.font.name = "Courier New"
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x1E, 0x1B, 0x4B)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                code_buffer = []
                in_code_block = False
            else:
                in_code_block = True
            continue

        if in_code_block:
            code_buffer.append(line)
            continue

        # 空行
        if not line.strip():
            continue

        # 标题
        if line.startswith("### "):
            p = doc.add_heading(line[4:], level=3)
            _set_heading_chinese(p)
            continue
        if line.startswith("## "):
            p = doc.add_heading(line[3:], level=2)
            _set_heading_chinese(p)
            continue
        if line.startswith("# "):
            p = doc.add_heading(line[2:], level=1)
            _set_heading_chinese(p)
            continue

        # 列表
        if line.strip().startswith("- ") or line.strip().startswith("* "):
            p = doc.add_paragraph(line.strip()[2:], style="List Bullet")
            _set_heading_chinese(p)
            continue

        # 序号列表
        match = re.match(r"^\d+[\.\、]\s*(.*)", line.strip())
        if match:
            p = doc.add_paragraph(match.group(1), style="List Number")
            _set_heading_chinese(p)
            continue

        # 普通段落（处理内联代码）
        parts = re.split(r"(`[^`]+`)", line)
        p = doc.add_paragraph()
        for part in parts:
            if part.startswith("`") and part.endswith("`"):
                run = p.add_run(part[1:-1])
                run.font.name = "Courier New"
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x43, 0x38, 0xCA)
            else:
                run = p.add_run(part)
                _set_chinese_font(run)

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()


# ─── PDF ───

def export_pdf(content: str) -> bytes:
    """将 Markdown 转换为 PDF 文件，正确处理中文。"""
    _register_fonts()

    buf = BytesIO()

    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
        leftMargin=2.5 * cm,
        rightMargin=2.5 * cm,
    )

    styles = getSampleStyleSheet()

    # 自定义样式 - 使用中文字体
    title_style = ParagraphStyle(
        "CustomTitle", parent=styles["Title"],
        fontName=_CN_FONT_BOLD, fontSize=18,
        textColor=colors.HexColor("#312E81"),
        spaceAfter=12,
    )
    heading1_style = ParagraphStyle(
        "CustomH1", parent=styles["Heading1"],
        fontName=_CN_FONT_BOLD, fontSize=16,
        textColor=colors.HexColor("#312E81"),
        spaceBefore=16, spaceAfter=8,
    )
    heading2_style = ParagraphStyle(
        "CustomH2", parent=styles["Heading2"],
        fontName=_CN_FONT_BOLD, fontSize=14,
        textColor=colors.HexColor("#312E81"),
        spaceBefore=12, spaceAfter=6,
    )
    heading3_style = ParagraphStyle(
        "CustomH3", parent=styles["Heading3"],
        fontName=_CN_FONT_BOLD, fontSize=12,
        textColor=colors.HexColor("#4338CA"),
        spaceBefore=8, spaceAfter=4,
    )
    body_style = ParagraphStyle(
        "CustomBody", parent=styles["Normal"],
        fontName=_CN_FONT, fontSize=10.5,
        leading=16, spaceAfter=6,
    )
    code_style = ParagraphStyle(
        "CustomCode", parent=styles["Code"],
        fontName="Courier", fontSize=8,
        leading=11, leftIndent=12,
        spaceBefore=4, spaceAfter=4,
        backColor=colors.HexColor("#F5F3FF"),
    )

    elements = []
    lines = content.split("\n")
    in_code_block = False
    code_buffer = []

    for line in lines:
        if line.startswith("```"):
            if in_code_block:
                code_text = "\n".join(code_buffer)
                # 转义 XML 特殊字符
                code_text = code_text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                elements.append(Preformatted(code_text, code_style))
                code_buffer = []
                in_code_block = False
            else:
                in_code_block = True
            continue

        if in_code_block:
            code_buffer.append(line)
            continue

        if not line.strip():
            elements.append(Spacer(1, 4))
            continue

        # 转义 XML 特殊字符
        safe_line = line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

        if line.startswith("### "):
            elements.append(Paragraph(safe_line[4:], heading3_style))
        elif line.startswith("## "):
            elements.append(Paragraph(safe_line[3:], heading2_style))
        elif line.startswith("# "):
            elements.append(Paragraph(safe_line[2:], heading1_style))
        elif line.strip().startswith("- ") or line.strip().startswith("* "):
            text = safe_line.strip()[2:]
            elements.append(Paragraph(f"• {text}", body_style))
        elif re.match(r"^\d+[\.\、]\s*(.*)", line.strip()):
            match = re.match(r"^\d+[\.\、]\s*(.*)", safe_line.strip())
            elements.append(Paragraph(match.group(1), body_style))
        else:
            # 内联代码处理
            text = re.sub(r"`([^`]+)`", r'<font face="Courier" color="#4338CA">\1</font>', safe_line)
            elements.append(Paragraph(text, body_style))

    doc.build(elements)
    buf.seek(0)
    return buf.getvalue()
